import io
import sys

data = sys.stdin.buffer.read()
kind = sys.argv[1]
try:
    if kind == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 80:
            raise ValueError('too many pages')
        text = '\n\n'.join((page.extract_text() or '') for page in reader.pages)
    elif kind == '.docx':
        from docx import Document
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text for cell in row.cells))
        text = '\n'.join(parts)
    else:
        raise ValueError('unsupported')
    sys.stdout.write(text[:30001])
except Exception:
    sys.exit(2)
